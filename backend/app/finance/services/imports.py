import hashlib,io
from datetime import date,datetime,timezone
from decimal import Decimal,InvalidOperation
from uuid import uuid4
from openpyxl import load_workbook
from ..domain.resources import FinanceRecordNotFound
from ..domain.errors import FinanceDomainError

class ImportValidationError(FinanceDomainError):
 status=422;code="VALIDATION_ERROR"
class ImportConflict(FinanceDomainError):
 status=409;code="STALE_VERSION"

HEADERS={"prices":("resourceCode","unitPrice","currency","effectiveFrom","scope"),"estimate":("resourceCode","activityExternalId","assignmentExternalId","originalQuantity","source")}
def parse_excel(content:bytes,kind:str,currency_unit:str|None=None):
 try:wb=load_workbook(io.BytesIO(content),read_only=True,data_only=True)
 except Exception as exc:raise ImportValidationError("invalid Excel workbook") from exc
 ws=wb.active
 rows=ws.iter_rows(values_only=True);header=tuple("" if x is None else str(x).strip() for x in next(rows,()))
 errors=[];normalized=[]
 missing=[x for x in HEADERS[kind] if x not in header]
 unexpected=[x for x in header if x and x not in HEADERS[kind]]
 if missing or unexpected:
  return [],([{"row":1,"field":x,"reason":"required_column"} for x in missing]
   +[{"row":1,"field":x,"reason":"unexpected_column"} for x in unexpected])
 positions={x:header.index(x) for x in HEADERS[kind]}
 for number,row in enumerate(rows,2):
  if not any(x is not None and str(x).strip() for x in row):continue
  item={k:(row[i] if i<len(row) else None) for k,i in positions.items()};item["rowNumber"]=number
  for key in ("resourceCode",):
   if item[key] is None or not str(item[key]).strip():errors.append({"row":number,"field":key,"reason":"required"})
   else:item[key]=str(item[key]).strip()
  for key in (("currency","scope") if kind=="prices" else ("activityExternalId","assignmentExternalId","source")):
   if item[key] is not None:item[key]=str(item[key]).strip()
  try:
   amount_key="unitPrice" if kind=="prices" else "originalQuantity";value=Decimal(str(item[amount_key]));item[amount_key]=format(value,"f")
   if value<0:errors.append({"row":number,"field":amount_key,"reason":"negative_value"})
   if kind=="prices" and item["currency"] not in ("IRR","TOMAN"):errors.append({"row":number,"field":"currency","reason":"invalid_choice"})
   if kind=="prices":
    irr=value if item["currency"]=="IRR" else value*10
    if irr!=irr.to_integral_value():errors.append({"row":number,"field":"unitPrice","reason":"fractional_irr"})
    if len(str(abs(irr.to_integral_value())))>18:errors.append({"row":number,"field":"unitPrice","reason":"max_digits"})
    item["unitPriceIrr"]=format(irr,".0f")
   elif value.as_tuple().exponent < -4:errors.append({"row":number,"field":"originalQuantity","reason":"max_decimal_places"})
  except (InvalidOperation,TypeError,ValueError):
   item[amount_key]=None;errors.append({"row":number,"field":amount_key,"reason":"invalid_decimal"})
  if kind=="prices":
   try:item["effectiveFrom"]=item["effectiveFrom"].isoformat() if hasattr(item["effectiveFrom"],"isoformat") else date.fromisoformat(str(item["effectiveFrom"])).isoformat()
   except (ValueError,TypeError):item["effectiveFrom"]=None;errors.append({"row":number,"field":"effectiveFrom","reason":"invalid_date"})
   if item["scope"] not in ("organization","project"):errors.append({"row":number,"field":"scope","reason":"invalid_choice"})
  elif item["source"] not in ("excel_import",):errors.append({"row":number,"field":"source","reason":"must_be_excel_import"})
  normalized.append(item)
 return normalized,errors

class FinanceImportService:
 def __init__(self,repo,id_factory=uuid4,clock=lambda:datetime.now(timezone.utc)):self.repo=repo;self.ids=id_factory;self.clock=clock
 async def preview(self,scope,kind,content,currency_unit=None):
  if scope.actor_user_id is None:raise PermissionError("authenticated actor is required")
  rows,errors=parse_excel(content,kind,currency_unit)
  resources=await self.repo.resolve_resources(scope,{str(row.get("resourceCode") or "").strip() for row in rows})
  resource_by_code={str(item["code"]):item for item in resources}
  for row in rows:
   code=str(row.get("resourceCode") or "").strip();resource=resource_by_code.get(code)
   if code and resource is None:errors.append({"row":row["rowNumber"],"field":"resourceCode","reason":"resource_not_found"})
   if resource is not None:
    row["resourceId"]=str(resource["id"]);row["resourceTitle"]=resource["title"];row["baseUnit"]=resource["base_unit"];row["resourceType"]=resource["resource_type"]
    if kind=="estimate" and resource["resource_type"]=="general_cost" and row.get("originalQuantity") is not None:
     amount=Decimal(row["originalQuantity"])
     if amount!=amount.to_integral_value():errors.append({"row":row["rowNumber"],"field":"originalQuantity","reason":"fractional_irr"})
     elif len(str(abs(amount.to_integral_value())))>18:errors.append({"row":row["rowNumber"],"field":"originalQuantity","reason":"max_digits"})
     else:row["originalUnitPriceIrr"]=format(amount,".0f")
  issues_by_row={row["rowNumber"]:[] for row in rows}
  for issue in errors:
   if issue["row"] in issues_by_row:issues_by_row[issue["row"]].append(issue)
  preview_rows=[]
  for row in rows:
   row_issues=issues_by_row[row["rowNumber"]]
   common={"rowNumber":row["rowNumber"],"status":"invalid" if row_issues else "valid","errors":row_issues,
    "resourceCode":row.get("resourceCode"),"resourceId":row.get("resourceId"),"resourceTitle":row.get("resourceTitle"),"baseUnit":row.get("baseUnit")}
   if kind=="estimate":common.update({"activityExternalId":row.get("activityExternalId"),"activityTitle":None,
    "assignmentExternalId":row.get("assignmentExternalId"),"originalQuantity":row.get("originalQuantity"),"source":row.get("source")})
   else:common.update({"unitPrice":row.get("unitPrice"),"currency":row.get("currency"),
    "normalizedUnitPriceIrr":row.get("unitPriceIrr"),"effectiveFrom":row.get("effectiveFrom"),"scope":row.get("scope")})
   preview_rows.append(common)
  valid_count=sum(not value["errors"] for value in preview_rows);invalid_count=len(preview_rows)-valid_count
  pid=self.ids();await self.repo.save_preview(scope,pid,kind,currency_unit,hashlib.sha256(content).hexdigest(),rows,errors,self.clock())
  return {"previewId":pid,"kind":kind,"rowCount":len(rows),"validCount":valid_count,"invalidCount":invalid_count,
   "rows":preview_rows,"errors":errors,"canCommit":not errors and bool(rows)}
 async def commit(self,scope,preview_id):
  try:result=await self.repo.commit(scope,preview_id,self.clock())
  except ValueError as exc:raise ImportConflict(str(exc)) from exc
  if result is None:raise FinanceRecordNotFound("import preview not found")
  return {"previewId":preview_id,"status":"committed","committedCount":result}
